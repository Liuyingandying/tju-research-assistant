"""科研调研报告导出器（v0.18 Phase 2.9-C-F）。

职责：将 ``ResearchReport`` 导出为 Markdown / Word / PDF 三种格式。

约束：

- **只读操作**：不修改 ``ResearchReport`` 对象，不调用 LLM，不读取论文库。
- **导出安全**：导出失败不影响原报告。
- **数据来源单一**：所有导出内容来自 ``ResearchReport`` 实例本身。
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import TYPE_CHECKING

from tju_info_retrieval.models.research_report import (
    NOT_GENERATED_TEXT,
    REPORT_SECTION_LABELS,
)

if TYPE_CHECKING:
    from tju_info_retrieval.models.research_report import ResearchReport


# ---------------------------------------------------------------------------
# 工具函数
# ---------------------------------------------------------------------------

def _section_number(key: str, all_keys: list[str]) -> str:
    """返回章节序号字符串，如 ``"1."``、``"2."``。"""
    try:
        idx = all_keys.index(key) + 1
    except ValueError:
        idx = 0
    return f"{idx}." if idx else ""


def _source_labels(
    report: "ResearchReport", paper_ids: list[str]
) -> list[str]:
    """将 paper_id 列表转换为 ``[P#]`` 编号列表。"""
    by_id: dict[str, "ReportSource"] = {
        p.paper_id: p for p in report.papers if p.paper_id
    }
    labels: list[str] = []
    for pid in paper_ids:
        paper = by_id.get(pid)
        if paper is not None:
            labels.append(f"[P{paper.index}]")
        else:
            labels.append(f"[?{pid[:8]}]")
    return labels


# ---------------------------------------------------------------------------
# Markdown 导出
# ---------------------------------------------------------------------------

def _build_markdown(report: "ResearchReport") -> str:
    """将 ``ResearchReport`` 渲染为 Markdown 字符串。"""
    lines: list[str] = []

    # --- 标题 ---
    lines.append(f"# {report.title}")
    lines.append("")

    # --- 报告信息 ---
    lines.append("## 报告信息")
    lines.append("")
    if report.created_time:
        lines.append(f"- 生成时间：{report.created_time[:19]}")
    lines.append(f"- 论文数量：{report.paper_count} 篇")
    provider = report.provider or "（未标注）"
    if report.model:
        provider = f"{provider} / {report.model}"
    lines.append(f"- Provider：{provider}")
    profile_ver = report.profile_version()
    lines.append(
        f"- 画像版本：{'p' + profile_ver if profile_ver else '未设置'}"
    )
    lines.append("")
    lines.append("---")
    lines.append("")

    # --- 章节 ---
    ordered = report.ordered_sections()
    all_keys = [s.key for s in ordered if s.key]

    for section in ordered:
        num = _section_number(section.key, all_keys)
        title = section.title or "（未命名章节）"
        lines.append(f"## {num}{title}")
        lines.append("")
        content = section.content or "（无正文）"
        lines.append(content)
        lines.append("")

        # 来源标注
        if section.source_papers:
            labels = _source_labels(report, section.source_papers)
            lines.append(f"来源：{' '.join(labels)}")
        else:
            lines.append("来源：（无）")

        if section.insufficient:
            lines.append("")
            lines.append(f"*说明：{NOT_GENERATED_TEXT}*")

        lines.append("")
        lines.append("---")
        lines.append("")

    # --- 参考论文 ---
    lines.append("## 参考论文")
    lines.append("")
    papers = report.papers
    if papers:
        for paper in papers:
            lines.append(f"[P{paper.index}] {paper.title}")
            lines.append("")
            if paper.authors:
                lines.append(f"- 作者：{'、'.join(paper.authors)}")
            if paper.year:
                lines.append(f"- 年份：{paper.year}")
            if paper.doi:
                lines.append(f"- DOI：{paper.doi}")
            if paper.source:
                lines.append(f"- 来源：{paper.source}")
            if paper.document_type:
                lines.append(f"- 文献类型：{paper.document_type}")
            lines.append("")
    else:
        lines.append("本报告没有来源论文快照。")
        lines.append("")

    return "\n".join(lines)


def export_markdown(report: "ResearchReport", path: str | Path) -> str:
    """导出报告为 Markdown 文件。

    Returns:
        实际写入的文件路径（字符串）。
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    content = _build_markdown(report)
    path.write_text(content, encoding="utf-8")
    return str(path)


# ---------------------------------------------------------------------------
# Word (DOCX) 导出
# ---------------------------------------------------------------------------

def _ensure_docx_import():
    """延迟导入 python-docx，避免无依赖时报错。"""
    try:
        from docx import Document  # noqa: F811
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        return Document, WD_ALIGN_PARAGRAPH, qn
    except ImportError as exc:
        raise ImportError(
            "导出 Word 需要安装 python-docx：pip install python-docx"
        ) from exc


def export_docx(report: "ResearchReport", path: str | Path) -> str:
    """导出报告为 Word (DOCX) 文件。

    Returns:
        实际写入的文件路径（字符串）。
    """
    Document, WD_ALIGN_PARAGRAPH, qn = _ensure_docx_import()

    doc = Document()

    # 设置中文字体
    style = doc.styles["Normal"]
    font = style.font
    font.name = "宋体"
    font._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    # --- 报告标题 ---
    title_para = doc.add_heading(report.title or "（未命名报告）", level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- 报告信息 ---
    meta_parts: list[str] = []
    if report.created_time:
        meta_parts.append(f"生成时间：{report.created_time[:19]}")
    meta_parts.append(f"论文数量：{report.paper_count} 篇")
    provider = report.provider or "（未标注）"
    if report.model:
        provider = f"{provider} / {report.model}"
    meta_parts.append(f"Provider：{provider}")
    profile_ver = report.profile_version()
    meta_parts.append(
        f"画像版本：{'p' + profile_ver if profile_ver else '未设置'}"
    )
    meta_para = doc.add_paragraph("  ｜  ".join(meta_parts))
    meta_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 分隔线
    doc.add_paragraph("_" * 40)

    # --- 章节 ---
    ordered = report.ordered_sections()
    all_keys = [s.key for s in ordered if s.key]

    for section in ordered:
        num = _section_number(section.key, all_keys)
        title = section.title or "（未命名章节）"
        doc.add_heading(f"{num}{title}", level=2)

        content = section.content or "（无正文）"
        doc.add_paragraph(content)

        # 来源标注
        if section.source_papers:
            labels = _source_labels(report, section.source_papers)
            src_para = doc.add_paragraph(f"来源：{' '.join(labels)}")
            src_para.runs[0].italic = True
        else:
            src_para = doc.add_paragraph("来源：（无）")
            src_para.runs[0].italic = True

        if section.insufficient:
            doc.add_paragraph(f"*说明：{NOT_GENERATED_TEXT}*")

        # 分隔线
        doc.add_paragraph("_" * 40)

    # --- 参考论文 ---
    doc.add_heading("参考论文", level=2)

    papers = report.papers
    if papers:
        for paper in papers:
            doc.add_heading(f"[P{paper.index}] {paper.title}", level=3)
            if paper.authors:
                doc.add_paragraph(f"作者：{'、'.join(paper.authors)}")
            if paper.year:
                doc.add_paragraph(f"年份：{paper.year}")
            if paper.doi:
                doc.add_paragraph(f"DOI：{paper.doi}")
            if paper.source:
                doc.add_paragraph(f"来源：{paper.source}")
            if paper.document_type:
                doc.add_paragraph(f"文献类型：{paper.document_type}")
    else:
        doc.add_paragraph("本报告没有来源论文快照。")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return str(path)


# ---------------------------------------------------------------------------
# PDF 导出
# ---------------------------------------------------------------------------

def _ensure_pdf_import():
    """延迟导入 reportlab，避免无依赖时报错。"""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm, cm
        from reportlab.platypus import (
            HRFlowable,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
        return (
            colors, A4, ParagraphStyle, getSampleStyleSheet,
            HRFlowable, Paragraph, SimpleDocTemplate,
            Spacer, Table, TableStyle, mm, cm,
        )
    except ImportError as exc:
        raise ImportError(
            "导出 PDF 需要安装 reportlab：pip install reportlab"
        ) from exc


def _find_chinese_font() -> str:
    """尝试找到可用的中文字体路径。"""
    # 常见中文字体路径
    font_paths = [
        # Windows
        os.path.join(os.environ.get("WINDIR", "C:\\Windows"),
                      "Fonts", "simsun.ttc"),
        os.path.join(os.environ.get("WINDIR", "C:\\Windows"),
                      "Fonts", "simhei.ttf"),
        os.path.join(os.environ.get("WINDIR", "C:\\Windows"),
                      "Fonts", "msyh.ttc"),
        # macOS
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Light.ttc",
        # Linux
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf",
    ]
    for fp in font_paths:
        if os.path.exists(fp):
            return fp
    return ""


def export_pdf(report: "ResearchReport", path: str | Path) -> str:
    """导出报告为 PDF 文件（支持中文）。

    Returns:
        实际写入的文件路径（字符串）。
    """
    (
        colors, A4, ParagraphStyle, getSampleStyleSheet,
        HRFlowable, Paragraph, SimpleDocTemplate,
        Spacer, Table, TableStyle, mm, cm,
    ) = _ensure_pdf_import()

    font_path = _find_chinese_font()

    page_w, page_h = A4
    margin = 2 * cm
    content_width = page_w - 2 * margin

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )

    styles = getSampleStyleSheet()

    # 注册中文字体
    if font_path:
        try:
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont

            # 尝试注册黑体（标题）
            try:
                pdfmetrics.registerFont(
                    TTFont("SimHei", font_path))
                pdfmetrics.registerFont(
                    TTFont("SimSun", font_path))
                bold_font = "SimHei"
                normal_font = "SimSun"
            except Exception:  # noqa: BLE001
                # 回退：用同一字体
                pdfmetrics.registerFont(
                    TTFont("ReportFont", font_path))
                bold_font = "ReportFont"
                normal_font = "ReportFont"
        except Exception:  # noqa: BLE001
            bold_font = "Helvetica-Bold"
            normal_font = "Helvetica"
    else:
        bold_font = "Helvetica-Bold"
        normal_font = "Helvetica"

    # 定义样式
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontName=bold_font,
        fontSize=18,
        spaceAfter=12,
        alignment=1,  # CENTER
    )

    meta_style = ParagraphStyle(
        "ReportMeta",
        parent=styles["Normal"],
        fontName=normal_font,
        fontSize=9,
        spaceAfter=6,
        alignment=1,  # CENTER
    )

    heading2_style = ParagraphStyle(
        "SectionHeading2",
        parent=styles["Heading2"],
        fontName=bold_font,
        fontSize=13,
        spaceBefore=10,
        spaceAfter=6,
    )

    body_style = ParagraphStyle(
        "ReportBody",
        parent=styles["Normal"],
        fontName=normal_font,
        fontSize=10,
        spaceAfter=6,
        leading=16,
    )

    source_style = ParagraphStyle(
        "SourceNote",
        parent=styles["Normal"],
        fontName=normal_font,
        fontSize=9,
        spaceAfter=4,
        leading=14,
        textColor=colors.darkgray,
    )

    warning_style = ParagraphStyle(
        "WarningNote",
        parent=styles["Normal"],
        fontName=normal_font,
        fontSize=9,
        spaceAfter=4,
        textColor=colors.orange,
        backColor="#FFF8E1",
        borderPadding=4,
    )

    ref_title_style = ParagraphStyle(
        "RefTitle",
        parent=styles["Heading3"],
        fontName=bold_font,
        fontSize=11,
        spaceBefore=6,
        spaceAfter=2,
    )

    ref_body_style = ParagraphStyle(
        "RefBody",
        parent=styles["Normal"],
        fontName=normal_font,
        fontSize=9,
        spaceAfter=2,
        leftIndent=12,
    )

    story = []

    # --- 报告标题 ---
    story.append(Paragraph(report.title or "（未命名报告）", title_style))

    # --- 报告信息 ---
    meta_parts: list[str] = []
    if report.created_time:
        meta_parts.append(f"生成时间：{report.created_time[:19]}")
    meta_parts.append(f"论文数量：{report.paper_count} 篇")
    provider = report.provider or "（未标注）"
    if report.model:
        provider = f"{provider} / {report.model}"
    meta_parts.append(f"Provider：{provider}")
    profile_ver = report.profile_version()
    meta_parts.append(
        f"画像版本：{'p' + profile_ver if profile_ver else '未设置'}"
    )
    story.append(Paragraph("  ｜  ".join(meta_parts), meta_style))

    # 分隔线
    story.append(
        HRFlowable(
            width="100%", thickness=1, color=colors.grey,
            spaceBefore=6, spaceAfter=6,
        )
    )

    # --- 章节 ---
    ordered = report.ordered_sections()
    all_keys = [s.key for s in ordered if s.key]

    for section in ordered:
        num = _section_number(section.key, all_keys)
        title = section.title or "（未命名章节）"
        story.append(Paragraph(f"{num}{title}", heading2_style))

        content = section.content or "（无正文）"
        story.append(Paragraph(content, body_style))

        # 来源标注
        if section.source_papers:
            labels = _source_labels(report, section.source_papers)
            story.append(
                Paragraph(f"来源：{' '.join(labels)}", source_style))
        else:
            story.append(Paragraph("来源：（无）", source_style))

        if section.insufficient:
            story.append(
                Paragraph(f"说明：{NOT_GENERATED_TEXT}", warning_style))

        # 分隔线
        story.append(
            HRFlowable(
                width="100%", thickness=0.5, color=colors.lightgrey,
                spaceBefore=6, spaceAfter=6,
            )
        )

    # --- 参考论文 ---
    story.append(Paragraph("参考论文", heading2_style))

    papers = report.papers
    if papers:
        for paper in papers:
            story.append(
                Paragraph(f"[P{paper.index}] {paper.title}", ref_title_style))
            if paper.authors:
                story.append(
                    Paragraph(
                        f"作者：{'、'.join(paper.authors)}", ref_body_style))
            if paper.year:
                story.append(
                    Paragraph(f"年份：{paper.year}", ref_body_style))
            if paper.doi:
                story.append(
                    Paragraph(f"DOI：{paper.doi}", ref_body_style))
            if paper.source:
                story.append(
                    Paragraph(f"来源：{paper.source}", ref_body_style))
            if paper.document_type:
                story.append(
                    Paragraph(
                        f"文献类型：{paper.document_type}", ref_body_style))
    else:
        story.append(Paragraph("本报告没有来源论文快照。", body_style))

    # 构建 PDF
    doc.build(story)
    return str(path)


# ---------------------------------------------------------------------------
# 公开 API
# ---------------------------------------------------------------------------

__all__ = [
    "export_markdown",
    "export_docx",
    "export_pdf",
    "ReportExporter",
]


class ReportExporter:
    """ResearchReport 导出器（只读，不修改报告）。

    所有导出方法均为静态方法，签名统一为：
    ``(report: ResearchReport, path: str | Path) -> str``
    返回实际写入的文件路径。
    """

    @staticmethod
    def export_markdown(report: "ResearchReport", path: str | Path) -> str:
        """导出为 Markdown 文件。"""
        return export_markdown(report, path)

    @staticmethod
    def export_docx(report: "ResearchReport", path: str | Path) -> str:
        """导出为 Word (DOCX) 文件。"""
        return export_docx(report, path)

    @staticmethod
    def export_pdf(report: "ResearchReport", path: str | Path) -> str:
        """导出为 PDF 文件（支持中文）。"""
        return export_pdf(report, path)
